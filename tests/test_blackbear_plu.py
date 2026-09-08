"""
tests/test_blackbear_plu.py
---------------------------
Offline unit tests for the Black Bear Burger PLU-code lookup (doc step 5) and
the sales calculation logic (doc step 4).

No browser, no credentials, no network required.

Run:
    /path/to/.venv/bin/python -m pytest tests/ -v
"""

import pathlib
import sys
from datetime import date

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).parent.parent))

import blackbear_convert as bb  # noqa: E402

MAPPING = {"black bear burger - victoria": "Black Bear Burger Victoria",
           "20ft fried chicken - victoria": "Black Bear Burger Victoria"}


def write_plu(path: pathlib.Path, rows, tab="PLU CODE"):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    ws = wb.create_sheet(tab)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


# ── PLU parsing ─────────────────────────────────────────────────────────


def test_plu_sheet_is_found_and_headers_sniffed(tmp_path):
    src = write_plu(tmp_path / "plu.xlsx", [
        ["Black Bear Burger PLU master", None],   # junk banner above the header
        ["Menu Item Name", "PLU Code"],
        ["Black Bear Burger", 5001],
        ["Cajun Fries", 5002],
    ])
    notes = []
    plu = bb.load_plu_codes(None, [src], notes)
    assert plu == {"black bear burger": "5001", "cajun fries": "5002"}
    assert any("header row 2" in n for n in notes)


def test_numeric_codes_lose_excel_float_suffix(tmp_path):
    src = write_plu(tmp_path / "plu.xlsx",
                    [["Item Name", "PLU"], ["Fries", 1234.0]])
    assert bb.load_plu_codes(None, [src], [])["fries"] == "1234"


def test_blank_names_and_codes_are_excluded(tmp_path):
    src = write_plu(tmp_path / "plu.xlsx", [
        ["Item Name", "PLU"],
        ["Fries", 10],
        [None, 11],          # blank name
        ["No Code Item", None],   # blank code
        [None, None],
    ])
    assert bb.load_plu_codes(None, [src], []) == {"fries": "10"}


def test_duplicate_item_keeps_first_code_and_warns(tmp_path):
    src = write_plu(tmp_path / "plu.xlsx", [
        ["Item Name", "PLU"],
        ["Fries", 10],
        ["  FRIES ", 99],    # same item, different case and code
    ])
    notes = []
    assert bb.load_plu_codes(None, [src], notes)["fries"] == "10"
    assert any("more than one PLU code" in n for n in notes)


def test_csv_plu_file_is_accepted(tmp_path):
    src = tmp_path / "plu.csv"
    src.write_text("Item Name,PLU Code\nFries,777\n", encoding="utf-8")
    assert bb.load_plu_codes(None, [src], [])["fries"] == "777"


def test_missing_plu_file_raises(tmp_path):
    try:
        bb.load_plu_codes(None, [tmp_path / "nope.xlsx"], [])
    except bb.ConvertError as exc:
        assert "PLU file not found" in str(exc)
    else:
        raise AssertionError("expected ConvertError")


def test_no_plu_source_returns_empty(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Items Sold 17th Aug"
    notes = []
    assert bb.load_plu_codes(wb, None, notes) == {}
    assert any("NONE" in n for n in notes)


# ── PLU applied to the Supy rows ────────────────────────────────────────


def _rows(plu, plu_missing):
    records = [{"restaurant": "Black Bear Burger - Victoria",
                "item": "Cajun Fries", "qty": 3, "gross": 12.00, "net": 10.00}]
    return bb.build_rows(records, date(2026, 8, 17), 0.20, MAPPING,
                         [], set(), plu, plu_missing)


def test_pos_item_id_uses_the_plu_code():
    missing = set()
    out = _rows({"cajun fries": "5002"}, missing)
    row = out["Black Bear Burger Victoria"][0]
    assert row["POS Item ID *"] == "5002"
    assert row["POS Item Name"] == "Cajun Fries"
    assert not missing


def test_pos_item_id_falls_back_to_name_and_is_flagged():
    missing = set()
    row = _rows({}, missing)["Black Bear Burger Victoria"][0]
    assert row["POS Item ID *"] == "Cajun Fries"
    assert missing == {"Cajun Fries"}


def test_plu_match_ignores_case_and_spacing():
    row = _rows({"cajun fries": "5002"}, set())["Black Bear Burger Victoria"][0]
    assert row["POS Item ID *"] == "5002"
    assert bb.norm_item("  Cajun   Fries. ") == "cajun fries"


# ── Sales calculation (doc step 4) ──────────────────────────────────────


def test_sales_values_follow_the_client_spec():
    """excl = after-discounts as-is; incl = excl x 1.2; discount = before - after."""
    row = _rows({}, set())["Black Bear Burger Victoria"][0]
    assert row["Sold QTY *"] == 3            # Count Orders (incl Undelivered)
    assert row["Total sales excl. tax *"] == 10.00
    assert row["Total sales incl. tax *"] == 12.00
    assert row["Total Discount Value"] == 2.00


def test_paired_restaurants_merge_into_one_supy_row():
    records = [
        {"restaurant": "Black Bear Burger - Victoria", "item": "Fries",
         "qty": 2, "gross": 10.00, "net": 8.00},
        {"restaurant": "20Ft Fried Chicken - Victoria", "item": "Fries",
         "qty": 1, "gross": 5.00, "net": 5.00},
    ]
    notes = []
    out = bb.build_rows(records, date(2026, 8, 17), 0.20, MAPPING,
                        notes, set(), {"fries": "10"}, set())
    rows = out["Black Bear Burger Victoria"]
    assert len(rows) == 1
    assert rows[0]["Sold QTY *"] == 3
    assert rows[0]["Total sales excl. tax *"] == 13.00
    assert rows[0]["Total Discount Value"] == 2.00
    assert notes and notes[0]["branch"] == "Black Bear Burger Victoria"


# ── The client's real workbook shape ────────────────────────────────────


def test_standard_plus_and_modifiers_plus_tabs_are_both_read(tmp_path):
    """'New Deliveroo - BBB.xlsx' keeps sku- codes and mod- codes on two tabs."""
    wb = openpyxl.Workbook()
    wb.active.title = "Victoria - 12th Oct"          # a sales tab, must be skipped
    wb.active.append(["Sales Date *", "POS Item ID *", "POS Item Name"])
    wb.active.append(["10/12/2025", "sku-999", "Should Not Be Used"])
    std = wb.create_sheet("Standard Plus")
    std.append(["EPoS Item Name", "EPoS Item ID"])
    std.append(["Meal Deal for 1", "sku-045"])
    std.append(["Black Bear", "sku-104"])
    mods = wb.create_sheet("Modifiers Plus")
    mods.append(["EPoS Item Name", "EPoS Item ID"])
    mods.append(["Standard Cooking", "mod-001"])
    src = tmp_path / "New Deliveroo - BBB.xlsx"
    wb.save(src)

    notes = []
    plu = bb.load_plu_codes(None, [src], notes)
    assert plu == {"meal deal for 1": "sku-045", "black bear": "sku-104",
                   "standard cooking": "mod-001"}
    assert "should not be used" not in plu


def test_sales_tab_named_like_a_plu_tab_is_skipped(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Standard Plus"
    ws.append(["Sales Date *", "POS Item ID *", "POS Item Name"])
    ws.append(["10/12/2025", "sku-999", "Fries"])
    src = tmp_path / "book.xlsx"
    wb.save(src)
    notes = []
    assert bb.load_plu_codes(None, [src], notes) == {}
    assert any("sales data, not a PLU master" in n for n in notes)


def test_two_masters_merge_and_conflicts_are_reported(tmp_path):
    a = write_plu(tmp_path / "bbb.xlsx",
                  [["Item Name", "PLU"], ["Fries", "sku-103"], ["Black Bear", "sku-104"]])
    b = write_plu(tmp_path / "20ft.xlsx",
                  [["Item Name", "PLU"], ["Cajun Fries", "sku-136"], ["Fries", "sku-999"]])
    notes = []
    plu = bb.load_plu_codes(None, [a, b], notes)
    assert plu == {"fries": "sku-103", "black bear": "sku-104",
                   "cajun fries": "sku-136"}          # first file wins
    assert any("keeping sku-103" in n for n in notes)
